"""Exportación a Excel (.xlsx) — estilo y armado ÚNICOS para toda la App.

Antes cada router repetía el Workbook + encabezado + Response, y el color del
encabezado estaba inconsistente (3E6086 vs 1F5FA8). Aquí vive la fuente única:
un solo color corporativo, un solo estilo de encabezado y un solo Response.
"""
from __future__ import annotations

from io import BytesIO

from fastapi import Response

# Color ÚNICO del encabezado de todas las exportaciones (azul corporativo).
XLSX_HDR = "1F5FA8"
_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def estilar_encabezado(ws, fila: int = 1) -> None:
    """Estilo estándar (blanco/negrita + relleno corporativo + centrado) a la fila
    de encabezados indicada (ya escrita)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor=XLSX_HDR)
    for cell in ws[fila]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def xlsx_response(wb, nombre: str) -> Response:
    """Serializa el Workbook y lo devuelve como descarga .xlsx."""
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(content=buf.read(), media_type=_MEDIA,
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


def tabla_simple(nombre: str, hoja: str, encabezados: list[str], filas,
                 anchos: list[int] | None = None) -> Response:
    """Hoja plana: encabezado estándar + filas + panel congelado. Para las
    exportaciones simples (claves, consultas…)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(encabezados)
    estilar_encabezado(ws, 1)
    for f in filas:
        ws.append(list(f))
    if anchos:
        for i, w in enumerate(anchos, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    return xlsx_response(wb, nombre)

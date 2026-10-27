"""Facturación mensual (histórico). Solo Super Usuarios.

Genera las filas del mes desde los servicios contratados (periodicidad Mes) de
los clientes activos, permite editarlas (valor, IVA, ReteIVA, observaciones) y
agregar filas manuales (Año/Eventual), guarda el histórico y exporta a Excel.
"""
from __future__ import annotations

import io
import json
import os
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel
from sqlalchemy import func
from sqlmodel import Session, select

from ..cuenta_cobro_docx import parse_cuenta_cobro_docx
from ..db import get_session
from ..deps import require_perm
from ..factura_xml import parse_factura_xml
from ..models import (Cliente, ClienteServicio, CuentaCobro, Emisor, Facturacion,
                      FacturacionItem, RolServicio, Usuario)
from ..scope import cliente_liviano

router = APIRouter(prefix="/api/facturacion", tags=["facturacion"])

MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def _dec(v) -> Decimal:
    try:
        return Decimal(str(v).replace(",", "").strip() or "0")
    except (InvalidOperation, AttributeError):
        return Decimal(0)


def _etiqueta(mes: int, anio: int) -> str:
    m = MESES[mes] if 1 <= mes <= 12 else str(mes)
    return f"{m} {anio}"


# ---------------------------------------------------------------- esquemas
class GenerarIn(BaseModel):
    anio: int
    mes: int


class ItemIn(BaseModel):
    cliente_id: int | None = None
    num_id: str = ""
    nombre: str = ""
    rol_servicio: str = ""
    concepto: str = ""
    valor: str = "0"
    forma: str = "Fa"
    iva: str = "0"
    reteiva: str = "0"
    periodicidad: str = "Mes"
    observaciones: str = ""


class GuardarIn(BaseModel):
    anio: int
    mes: int
    items: list[ItemIn]


def _item_dict(it: FacturacionItem) -> dict:
    return {"id": it.id, "cliente_id": it.cliente_id, "num_id": it.num_id,
            "nombre": it.nombre, "rol_servicio": it.rol_servicio, "concepto": it.concepto,
            "valor": str(it.valor or 0), "forma": it.forma, "iva": str(it.iva or 0),
            "reteiva": str(it.reteiva or 0), "periodicidad": it.periodicidad,
            "observaciones": it.observaciones}


# ---------------------------------------------------------------- generar (vista previa)
@router.post("/generar")
def generar(data: GenerarIn, session: Session = Depends(get_session),
            user: Usuario = Depends(require_perm("facturacion.crear"))):
    """Arma las filas del período desde los cobros programados de los clientes
    activos: cada servicio tiene un ``periodo`` (AAAA-MM) y se toman los que
    coinciden con el período pedido y aún están PENDIENTES (facturacion_id NULL).
    NO guarda. IVA por defecto 19% si Fa (0 si CC); ReteIVA 0."""
    etiqueta = _etiqueta(data.mes, data.anio)
    periodo = f"{data.anio:04d}-{data.mes:02d}"
    from ..scope import confid_ocultos as _confid_ocultos
    ocultos = _confid_ocultos(session, user)      # confidenciales no visibles para este usuario
    roles = {r.id: r.descripcion for r in session.exec(select(RolServicio)).all()}
    servicios = session.exec(select(ClienteServicio).where(
        ClienteServicio.periodo == periodo,
        ClienteServicio.facturacion_id == None)).all()   # noqa: E711
    # Solo los clientes que aparecen en el período (no toda la tabla cliente).
    cli_ids = {sv.cliente_id for sv in servicios}
    clientes = ({c.id: c for c in session.exec(
                    select(Cliente).options(*cliente_liviano()).where(Cliente.id.in_(cli_ids)))}
                if cli_ids else {})
    filas = []
    for sv in servicios:
        c = clientes.get(sv.cliente_id)
        if not c or c.estado != "Act":
            continue
        if c.id in ocultos:                        # confidencial: supervisor O asignado
            continue
        rol = roles.get(sv.rol_servicio_id, "")
        valor = sv.valor_servicio or Decimal(0)
        forma = sv.forma_facturacion or "Fa"
        iva = (valor * Decimal("0.19")).quantize(Decimal("1")) if forma == "Fa" else Decimal(0)
        filas.append({
            "cliente_id": c.id, "num_id": c.num_id, "nombre": c.nombre,
            "rol_servicio": rol, "concepto": f"{rol} — {etiqueta}".strip(" —"),
            "valor": str(int(valor)), "forma": forma, "iva": str(int(iva)),
            "reteiva": "0", "periodicidad": "Mes", "observaciones": "",
        })
    filas.sort(key=lambda x: x["nombre"].lower())
    return {"etiqueta": etiqueta, "items": filas}


# ---------------------------------------------------------------- servicios no mensuales
@router.get("/no-mensuales")
def no_mensuales(_u: Usuario = Depends(require_perm("facturacion.ver"))):
    """(Obsoleto) Con el modelo de cobros por período ya no hay servicios
    Año/Eventual: todo cobro se programa con su ``periodo`` y sale al Generar.
    Se conserva la ruta devolviendo lista vacía para no romper el frontend."""
    return []


# ---------------------------------------------------------------- guardar
@router.post("", status_code=201)
def guardar(data: GuardarIn, session: Session = Depends(get_session),
            user: Usuario = Depends(require_perm("facturacion.crear"))):
    # Una sola facturación por mes: si el período ya existe, se reemplaza.
    previas = session.exec(select(Facturacion).where(
        Facturacion.anio == data.anio, Facturacion.mes == data.mes)).all()
    prev_ids = [p.id for p in previas]
    if prev_ids:   # al reemplazar, los cobros marcados por la previa vuelven a pendiente
        for sv in session.exec(select(ClienteServicio).where(
                ClienteServicio.facturacion_id.in_(prev_ids))).all():
            sv.facturacion_id = None
            session.add(sv)
        session.flush()
    for prev in previas:
        for it in session.exec(select(FacturacionItem).where(
                FacturacionItem.facturacion_id == prev.id)).all():
            session.delete(it)
        session.flush()
        session.delete(prev)
    if previas:
        session.commit()

    cab = Facturacion(anio=data.anio, mes=data.mes, etiqueta=_etiqueta(data.mes, data.anio),
                      creado_por=user.usuario or user.nombre)
    session.add(cab)
    session.commit()
    session.refresh(cab)
    for i, it in enumerate(data.items):
        session.add(FacturacionItem(
            facturacion_id=cab.id, cliente_id=it.cliente_id, num_id=it.num_id,
            nombre=it.nombre, rol_servicio=it.rol_servicio, concepto=it.concepto,
            valor=_dec(it.valor), forma=it.forma, iva=_dec(it.iva),
            reteiva=_dec(it.reteiva), periodicidad=it.periodicidad,
            observaciones=it.observaciones, orden=i))
    session.commit()

    # Marca como facturados los cobros programados de ese período de los clientes
    # incluidos (aún pendientes) → dejan de salir al Generar y muestran el estado.
    periodo = f"{data.anio:04d}-{data.mes:02d}"
    cli_ids = {it.cliente_id for it in data.items if it.cliente_id}
    if cli_ids:
        for sv in session.exec(select(ClienteServicio).where(
                ClienteServicio.periodo == periodo,
                ClienteServicio.cliente_id.in_(cli_ids),
                ClienteServicio.facturacion_id == None)).all():   # noqa: E711
            sv.facturacion_id = cab.id
            session.add(sv)
        session.commit()
    return {"id": cab.id, "etiqueta": cab.etiqueta, "items": len(data.items),
            "reemplazo": bool(previas)}


# ---------------------------------------------------------------- histórico
@router.get("")
def listar(session: Session = Depends(get_session), _u: Usuario = Depends(require_perm("facturacion.ver"))):
    cabs = session.exec(select(Facturacion).order_by(Facturacion.anio.desc(),
                                                     Facturacion.mes.desc(),
                                                     Facturacion.id.desc())).all()
    # Conteo y total por cabecera agregados en la BD (no se traen todos los ítems).
    agg = session.exec(
        select(FacturacionItem.facturacion_id, func.count(),
               func.coalesce(func.sum(FacturacionItem.valor), 0))
        .group_by(FacturacionItem.facturacion_id)).all()
    conteo: dict[int, int] = {fid: n for fid, n, _ in agg}
    tot: dict[int, Decimal] = {fid: Decimal(str(s)) for fid, _, s in agg}
    return [{"id": c.id, "anio": c.anio, "mes": c.mes, "etiqueta": c.etiqueta,
             "creado_at": c.creado_at.isoformat(), "creado_por": c.creado_por,
             "aprobada_at": c.aprobada_at.isoformat() if c.aprobada_at else None,
             "aprobada_por": c.aprobada_por,
             "n_items": conteo.get(c.id, 0), "total_valor": str(tot.get(c.id, Decimal(0)))}
            for c in cabs]


# ---------------------------------------------------------- validar vs emitidas
_ndig = lambda s: re.sub(r"\D", "", str(s or ""))   # solo dígitos (NIT sin DV/puntos)


class ValidarIn(BaseModel):
    carpeta: str


@router.post("/{fid}/validar")
def validar(fid: int, data: ValidarIn, session: Session = Depends(get_session),
            _u: Usuario = Depends(require_perm("facturacion.validar"))):
    """Cruza la facturación aprobada contra las facturas electrónicas emitidas
    (XML UBL) que estén en la CARPETA indicada (se lee completa, incluyendo
    subcarpetas). Compara por cliente (NIT): base gravable y IVA. Las CC no
    tienen XML → se listan aparte. Las retenciones no vienen en la venta."""
    cab = session.get(Facturacion, fid)
    if not cab:
        raise HTTPException(404, "Facturación no encontrada")
    items = session.exec(select(FacturacionItem).where(
        FacturacionItem.facturacion_id == fid)).all()

    # App local (el backend corre en la máquina del propio usuario): se permite
    # cualquier carpeta del equipo. Solo se valida que exista y sea un directorio.
    try:
        carpeta = Path(data.carpeta).resolve()
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "Ruta inválida")
    if not carpeta.is_dir():
        raise HTTPException(400, "La carpeta no existe")

    # 1) Leer los XML de la carpeta (recursivo, dedupe por número)
    facturas: list[dict] = []
    errores: list[dict] = []
    vistos: set[str] = set()
    n_xml = 0
    for p in sorted(carpeta.rglob("*.xml")):
        n_xml += 1
        try:
            raw = p.read_bytes()
        except OSError:
            continue
        try:
            fx = parse_factura_xml(raw)
        except Exception as exc:  # noqa: BLE001
            errores.append({"archivo": p.name, "error": str(exc)[:120]})
            continue
        if not fx.get("numero"):
            continue                       # XML que no es factura (ej. acuse) → se ignora
        if fx["numero"] in vistos:
            continue
        vistos.add(fx["numero"])
        facturas.append(fx)

    # 1b) Leer las cuentas de cobro (.docx) de la carpeta (recursivo)
    cuentas: list[dict] = []
    vistas: set[str] = set()
    n_docx = 0
    for p in sorted(carpeta.rglob("*.docx")):
        if p.name.startswith("~$"):
            continue                          # archivo temporal de Word bloqueado
        n_docx += 1
        try:
            raw = p.read_bytes()
        except OSError:
            continue
        try:
            cc = parse_cuenta_cobro_docx(raw)
        except Exception as exc:  # noqa: BLE001
            errores.append({"archivo": p.name, "error": str(exc)[:120]})
            continue
        if not cc.get("cliente_nit") or cc.get("valor", 0) <= 0:
            continue                          # .docx que no es cuenta de cobro
        clave = cc.get("numero") or p.name
        if clave in vistas:
            continue
        vistas.add(clave)
        cuentas.append(cc)

    # 2) Esperado agrupado por NIT, separando Fa y CC según la facturación
    esp_fa: dict[str, dict] = {}
    esp_cc: dict[str, dict] = {}
    for it in items:
        destino = esp_cc if (it.forma or "Fa") == "CC" else esp_fa
        k = _ndig(it.num_id)
        e = destino.setdefault(k, {"nombre": it.nombre, "base": 0.0, "iva": 0.0})
        e["base"] += float(it.valor or 0)
        e["iva"] += float(it.iva or 0)

    # 3) Emitido agrupado por NIT (Fa desde XML, CC desde docx)
    emi_fa: dict[str, dict] = {}
    for fx in facturas:
        k = _ndig(fx["cliente_nit"])
        m = emi_fa.setdefault(k, {"nombre": fx["cliente_nombre"], "base": 0.0, "iva": 0.0, "numeros": []})
        m["base"] += fx["base"]; m["iva"] += fx["iva"]; m["numeros"].append(fx["numero"])
    emi_cc: dict[str, dict] = {}
    for cc in cuentas:
        k = _ndig(cc["cliente_nit"])
        m = emi_cc.setdefault(k, {"nombre": cc["cliente_nombre"], "base": 0.0, "iva": 0.0, "numeros": []})
        m["base"] += cc["valor"]; m["numeros"].append(cc["numero"])

    def _cruzar(esperado: dict, emitido: dict, con_iva: bool) -> list[dict]:
        out = []
        for k in sorted(set(esperado) | set(emitido)):
            e, m = esperado.get(k), emitido.get(k)
            base_esp = e["base"] if e else 0.0
            iva_esp = e["iva"] if (e and con_iva) else 0.0
            base_emi = m["base"] if m else 0.0
            iva_emi = m["iva"] if (m and con_iva) else 0.0
            nombre = (e or m)["nombre"]
            db, di = base_emi - base_esp, iva_emi - iva_esp
            if e and m:
                estado = "ok" if abs(db) < 1 and abs(di) < 1 else "dif"
            elif e:
                estado = "sin_emitir"
            else:
                estado = "sin_facturar"
            out.append({"nit": k, "nombre": nombre, "estado": estado,
                        "base_esp": base_esp, "base_emi": base_emi, "dif_base": db,
                        "iva_esp": iva_esp, "iva_emi": iva_emi, "dif_iva": di,
                        "numeros": (m["numeros"] if m else [])})
        return out

    filas = _cruzar(esp_fa, emi_fa, con_iva=True)
    filas_cc = _cruzar(esp_cc, emi_cc, con_iva=False)

    def _cnt(fs, est):
        return sum(1 for f in fs if f["estado"] == est)

    resumen = {
        "n_xml": n_xml, "n_facturas": len(facturas),
        "n_docx": n_docx, "n_cc": len(cuentas),
        "ok": _cnt(filas, "ok"), "dif": _cnt(filas, "dif"),
        "sin_emitir": _cnt(filas, "sin_emitir"), "sin_facturar": _cnt(filas, "sin_facturar"),
        "cc_ok": _cnt(filas_cc, "ok"), "cc_dif": _cnt(filas_cc, "dif"),
        "cc_sin_emitir": _cnt(filas_cc, "sin_emitir"), "cc_sin_facturar": _cnt(filas_cc, "sin_facturar"),
        "base_esp": sum(f["base_esp"] for f in filas),
        "base_emi": sum(f["base_emi"] for f in filas),
        "iva_esp": sum(f["iva_esp"] for f in filas),
        "iva_emi": sum(f["iva_emi"] for f in filas),
        "cc_base_esp": sum(f["base_esp"] for f in filas_cc),
        "cc_base_emi": sum(f["base_emi"] for f in filas_cc),
    }
    return {"etiqueta": cab.etiqueta, "resumen": resumen,
            "filas": filas, "filas_cc": filas_cc, "errores": errores}


# ---------------------------------------------------------------- aprobar
@router.post("/{fid}/aprobar")
def aprobar(fid: int, session: Session = Depends(get_session),
            user: Usuario = Depends(require_perm("facturacion.aprobar"))):
    """Aprueba (o desaprueba) la facturación del mes. Guarda fecha + usuario.
    Solo superusuarios. Base para validar luego contra las facturas emitidas."""
    cab = session.get(Facturacion, fid)
    if not cab:
        raise HTTPException(404, "Facturación no encontrada")
    if cab.aprobada_at:                       # ya aprobada → desaprobar (toggle)
        cab.aprobada_at = None
        cab.aprobada_por = ""
    else:
        cab.aprobada_at = datetime.utcnow()
        cab.aprobada_por = user.usuario or user.nombre
    session.add(cab)
    session.commit()
    return {"aprobada_at": cab.aprobada_at.isoformat() if cab.aprobada_at else None,
            "aprobada_por": cab.aprobada_por}


@router.get("/{fid}")
def obtener(fid: int, session: Session = Depends(get_session),
            user: Usuario = Depends(require_perm("facturacion.ver"))):
    cab = session.get(Facturacion, fid)
    if not cab:
        raise HTTPException(404, "Facturación no encontrada")
    items = session.exec(select(FacturacionItem).where(FacturacionItem.facturacion_id == fid)
                         .order_by(FacturacionItem.orden)).all()
    from ..scope import confid_ocultos as _confid_ocultos
    ocultos = _confid_ocultos(session, user)      # oculta ítems de confidenciales no visibles
    if ocultos:
        items = [i for i in items if i.cliente_id not in ocultos]
    return {"id": cab.id, "anio": cab.anio, "mes": cab.mes, "etiqueta": cab.etiqueta,
            "creado_at": cab.creado_at.isoformat(), "creado_por": cab.creado_por,
            "items": [_item_dict(i) for i in items]}


@router.delete("/{fid}", status_code=204)
def borrar(fid: int, session: Session = Depends(get_session),
           _u: Usuario = Depends(require_perm("facturacion.eliminar"))):
    cab = session.get(Facturacion, fid)
    if not cab:
        raise HTTPException(404, "Facturación no encontrada")
    for sv in session.exec(select(ClienteServicio).where(
            ClienteServicio.facturacion_id == fid)).all():   # los cobros vuelven a pendiente
        sv.facturacion_id = None
        session.add(sv)
    for it in session.exec(select(FacturacionItem).where(FacturacionItem.facturacion_id == fid)).all():
        session.delete(it)
    session.flush()          # borra los ítems antes que el encabezado (evita violar la FK)
    session.delete(cab)
    session.commit()


# ---------------------------------------------------------------- exportar Excel
def _excel_response(etiqueta: str, filas: list) -> Response:
    """filas: lista de objetos/dicts con num_id, nombre, concepto, valor, forma,
    iva, reteiva, periodicidad, observaciones. Devuelve el xlsx como Response."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from ..export_xlsx import estilar_encabezado, xlsx_response

    def g(f, k):
        return getattr(f, k) if not isinstance(f, dict) else f.get(k)

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturacion"
    encab = ["Id", "Nombre Tercero", "Concepto", "Valor", "Factura", "IVA",
             "ReteIVA", "Periodicidad", "Observaciones"]
    ws.append([f"Facturación {etiqueta}"])
    ws.append([])
    ws.append(encab)
    estilar_encabezado(ws, 3)
    for f in filas:
        ws.append([g(f, "num_id"), g(f, "nombre"), g(f, "concepto"),
                   float(_dec(g(f, "valor"))), g(f, "forma"),
                   float(_dec(g(f, "iva"))), float(_dec(g(f, "reteiva"))),
                   g(f, "periodicidad"), g(f, "observaciones")])
    n = len(filas)
    if n:
        fila_tot = 3 + n + 1
        ws.cell(row=fila_tot, column=3, value="TOTALES").font = Font(bold=True)
        for col in (4, 6, 7):
            letra = ws.cell(row=4, column=col).column_letter
            ws.cell(row=fila_tot, column=col,
                    value=f"=SUM({letra}4:{letra}{3+n})").font = Font(bold=True)
    anchos = [12, 34, 40, 14, 9, 13, 13, 12, 30]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = w
    for col in ("D", "F", "G"):
        for r in range(4, 4 + n):
            ws[f"{col}{r}"].number_format = "#,##0"

    nombre = f"Facturacion {etiqueta}.xlsx".replace(" ", "_")
    return xlsx_response(wb, nombre)


@router.post("/exportar")
def exportar_grilla(data: GuardarIn, _u: Usuario = Depends(require_perm("facturacion.exportar"))):
    """Exporta a Excel las filas actuales (sin guardar)."""
    return _excel_response(_etiqueta(data.mes, data.anio), data.items)


@router.get("/{fid}/excel")
def exportar(fid: int, session: Session = Depends(get_session),
             user: Usuario = Depends(require_perm("facturacion.exportar"))):
    cab = session.get(Facturacion, fid)
    if not cab:
        raise HTTPException(404, "Facturación no encontrada")
    items = session.exec(select(FacturacionItem).where(FacturacionItem.facturacion_id == fid)
                         .order_by(FacturacionItem.orden)).all()
    from ..scope import confid_ocultos as _confid_ocultos
    ocultos = _confid_ocultos(session, user)
    if ocultos:
        items = [i for i in items if i.cliente_id not in ocultos]
    return _excel_response(cab.etiqueta, items)


class ValidarExcelIn(BaseModel):
    etiqueta: str = ""
    resumen: dict = {}
    filas: list[dict] = []
    filas_cc: list[dict] = []


_EST_LBL = {"ok": "Coincide", "dif": "Diferencia",
            "sin_emitir": "Sin emitir", "sin_facturar": "Sin facturar"}


@router.post("/validar/excel")
def validar_excel(data: ValidarExcelIn, _u: Usuario = Depends(require_perm("facturacion.validar"))):
    """Exporta a Excel el reporte de validación (facturas + cuentas de cobro)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from ..export_xlsx import estilar_encabezado, xlsx_response

    wb = Workbook()
    ws = wb.active
    ws.title = "Validación"

    ws["A1"] = f"Validación facturación {data.etiqueta} — aprobada vs emitidas"
    ws["A1"].font = Font(bold=True, size=13)
    R = data.resumen or {}
    ws["A2"] = (f"XML: {R.get('n_xml',0)} → Facturas: {R.get('n_facturas',0)}   |   "
                f"Word: {R.get('n_docx',0)} → CC: {R.get('n_cc',0)}   |   "
                f"Coinciden: {R.get('ok',0)+R.get('cc_ok',0)}   Diferencias: {R.get('dif',0)+R.get('cc_dif',0)}   "
                f"Sin emitir: {R.get('sin_emitir',0)+R.get('cc_sin_emitir',0)}   "
                f"Sin facturar: {R.get('sin_facturar',0)+R.get('cc_sin_facturar',0)}")
    ws["A2"].font = Font(color="555555")

    # Una sola tabla: facturas (Fa) + cuentas de cobro (CC), ordenada por NIT.
    todas = ([{**f, "tipo": "Fa"} for f in data.filas]
             + [{**f, "tipo": "CC"} for f in data.filas_cc])
    todas.sort(key=lambda f: str(f.get("nit") or ""))

    cols = ["NIT", "Cliente", "Valor aprob.", "Valor emit.", "Dif",
            "IVA aprob.", "IVA emit.", "Dif IVA", "Estado", "Fact/CC"]
    fila = 4
    for j, c in enumerate(cols, 1):
        ws.cell(row=fila, column=j, value=c)
    estilar_encabezado(ws, fila)
    fila += 1
    for f in todas:
        docs = f"{f.get('tipo')} " + ", ".join(f.get("numeros") or [])
        vals = [f.get("nit"), f.get("nombre"), f.get("base_esp"), f.get("base_emi"), f.get("dif_base"),
                f.get("iva_esp"), f.get("iva_emi"), f.get("dif_iva"),
                _EST_LBL.get(f.get("estado"), f.get("estado")), docs.strip()]
        for j, v in enumerate(vals, 1):
            ws.cell(row=fila, column=j, value=v)
        fila += 1

    for col in ("C", "D", "E", "F", "G", "H"):
        for r in range(5, fila):
            ws[f"{col}{r}"].number_format = "#,##0"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 42
    for c in ("C", "D", "E", "F", "G", "H"):
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["I"].width = 13
    ws.column_dimensions["J"].width = 26

    nombre = f"Validacion {data.etiqueta}.xlsx".replace(" ", "_")
    return xlsx_response(wb, nombre)


# ===================================================== Cuentas de cobro (CC)
def _cc_validar_cliente(session: Session, fid: int, cliente_id: int):
    """Valida que el cliente tenga ítems CC en la facturación y un emisor asignado.
    Devuelve (items_cc, cliente, emisor, deudor_doc). La CC agrupa TODOS los
    conceptos CC del cliente en el mes. El deudor va como NIT + dígito de
    verificación aunque el cliente esté registrado como cédula (lo pide el formato)."""
    from ..util import calcular_dv
    items = session.exec(select(FacturacionItem).where(
        FacturacionItem.facturacion_id == fid,
        FacturacionItem.cliente_id == cliente_id)).all()
    items_cc = [it for it in items if (it.forma or "").upper() == "CC"]
    if not items_cc:
        raise HTTPException(404, "El cliente no tiene cuentas de cobro (ítems CC) en esta facturación.")
    cli = session.get(Cliente, cliente_id)
    if not cli or not cli.emisor_id:
        raise HTTPException(400, "El cliente no tiene un emisor asignado (ficha → Comercial).")
    em = session.get(Emisor, cli.emisor_id)
    if not em:
        raise HTTPException(400, "El emisor asignado ya no existe.")
    num_fmt = f"{int(cli.num_id):,}".replace(",", ".") if (cli.num_id or "").isdigit() else (cli.num_id or "")
    dv = cli.dv if cli.dv is not None else calcular_dv(cli.num_id)
    deudor_doc = f"NIT. {num_fmt}" + (f" - {dv}" if dv is not None else "")
    return items_cc, cli, em, deudor_doc


def _conceptos_de(items_cc) -> "tuple[list, Decimal]":
    """Renglones [{"concepto","valor"}] + total, ordenados por valor desc. Reúne los
    ítems CC del cliente en la lista que viaja a la CC y al PDF."""
    filas = sorted(items_cc, key=lambda it: float(it.valor or 0), reverse=True)
    conceptos = [{"concepto": it.concepto or "", "valor": float(it.valor or 0)} for it in filas]
    total = sum((it.valor or Decimal(0)) for it in filas) or Decimal(0)
    return conceptos, total


def _cc_existente(session: Session, anio: int, mes: int, cliente_id) -> "CuentaCobro | None":
    """La cuenta de cobro EN FIRME ya generada para (año, mes, cliente): UNA por
    cliente/mes, agrupando todos sus conceptos CC. Identidad de negocio estable
    (sobrevive al re-guardado de la facturación)."""
    return session.exec(select(CuentaCobro).where(
        CuentaCobro.anio == anio, CuentaCobro.mes == mes,
        CuentaCobro.cliente_id == cliente_id).order_by(CuentaCobro.id)).first()


def _cc_html(session: Session, cc: "CuentaCobro") -> str:
    """Reconstruye el HTML del documento desde la FOTO guardada en la cuenta de cobro.
    La firma cae a la del emisor ACTUAL si la foto quedó vacía (CC generadas antes de
    cargarle la firma al emisor); las que ya la tienen conservan la congelada."""
    from ..cuenta_cobro import construir_html
    firma = cc.emisor_firma
    if not firma and cc.emisor_id:
        em = session.get(Emisor, cc.emisor_id)
        firma = em.firma if em else ""
    # Renglones de concepto: de la foto JSON; CC viejas (sin `conceptos`) = un renglón.
    try:
        conceptos = json.loads(cc.conceptos) if cc.conceptos else []
    except (ValueError, TypeError):
        conceptos = []
    if not conceptos:
        conceptos = [{"concepto": cc.concepto, "valor": float(cc.valor or 0)}]
    return construir_html(
        numero=cc.numero, fecha=cc.fecha, deudor=cc.deudor, deudor_doc=cc.deudor_doc,
        emisor_nombre=cc.emisor_nombre, emisor_cedula=cc.emisor_cedula,
        emisor_direccion=cc.emisor_direccion, emisor_celular=cc.emisor_celular,
        nota_383=cc.nota_383, nota_383_texto=cc.nota_383_texto, firma=firma,
        conceptos=conceptos, valor=cc.valor or 0,
        pago_banco=cc.pago_banco, pago_tipo=cc.pago_tipo,
        pago_numero=cc.pago_numero, pago_titular=cc.pago_titular)


def _slug_archivo(s: str) -> str:
    """Nombre de archivo seguro (sin caracteres inválidos de Windows)."""
    return re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", (s or "")).strip()[:70] or "cliente"


def _ruta_cc(carpeta: str, numero: str, cliente: str, anio: int, mes: int) -> str:
    """Ruta destino del PDF dentro de la carpeta base (organizada por año):
    «<carpeta>\\<año>\\<N°> - CC - <cliente> - <AAAA MM>.pdf». '' si no hay carpeta."""
    if not (carpeta or "").strip():
        return ""
    nombre = f"{numero} - CC - {_slug_archivo(cliente)} - {anio} {mes:02d}.pdf"
    return os.path.join(carpeta.strip(), str(anio), nombre)


def _generar_cc(session: Session, fid: int, cliente_id: int, creado_por: str,
                pago: dict | None = None) -> "tuple[CuentaCobro, bool]":
    """Genera la cuenta de cobro EN FIRME del CLIENTE (una por año/mes/cliente,
    agrupando todos sus conceptos CC). Consume el consecutivo del emisor y guarda
    la foto. Idempotente por la identidad de negocio: si ya existe, la devuelve sin
    quemar otro número. Devuelve (cuenta, nueva?). `pago` se puede pasar
    precalculado (generación en lote) para no releer los parámetros por cliente."""
    from datetime import date
    from .parametros import _cobro
    if pago is None:
        pago = _cobro(session)
    fac = session.get(Facturacion, fid)
    if not fac:
        raise HTTPException(404, "Facturación no encontrada")
    items_cc, cli, em, deudor_doc = _cc_validar_cliente(session, fid, cliente_id)
    ya = _cc_existente(session, fac.anio, fac.mes, cli.id)
    if ya:
        if ya.facturacion_id != fid:   # mantiene al día el enlace informativo
            ya.facturacion_id = fid
            session.add(ya)
            session.commit()
            session.refresh(ya)
        return ya, False
    conceptos, total = _conceptos_de(items_cc)
    concepto_lbl = " · ".join(c["concepto"] for c in conceptos if c["concepto"])
    # Bloquea la fila del emisor para que dos generaciones no tomen el mismo N°.
    em = session.exec(select(Emisor).where(Emisor.id == em.id).with_for_update()).one()
    nuevo = em.consecutivo + 1
    em.consecutivo = nuevo
    numero = f"{em.prefijo}-{nuevo}" if em.prefijo else str(nuevo)
    cc = CuentaCobro(
        anio=fac.anio, mes=fac.mes, concepto=concepto_lbl, conceptos=json.dumps(conceptos, ensure_ascii=False),
        facturacion_id=fid, emisor_id=em.id, cliente_id=cli.id,
        numero=numero, prefijo=em.prefijo, consecutivo=nuevo, fecha=date.today(),
        deudor=cli.nombre, deudor_doc=deudor_doc,
        email_destino=(cli.email or "").strip(),
        emisor_nombre=em.nombre, emisor_cedula=em.cedula, emisor_direccion=em.direccion,
        emisor_celular=em.celular, nota_383=em.nota_383, nota_383_texto=pago["nota_383"],
        emisor_firma=em.firma,
        valor=total,
        pago_banco=pago["banco"], pago_tipo=pago["tipo_cuenta"],
        pago_numero=pago["numero_cuenta"], pago_titular=pago["titular"],
        estado="generada", creado_por=creado_por)
    session.add(em)
    session.add(cc)
    session.commit()
    session.refresh(cc)
    # Guarda el PDF en la carpeta configurada (si hay). No rompe la generación si falla.
    ruta = _ruta_cc(pago.get("carpeta", ""), numero, cc.deudor, cc.anio, cc.mes)
    if ruta:
        try:
            from ..cuenta_cobro import html_a_pdf
            os.makedirs(os.path.dirname(ruta), exist_ok=True)
            with open(ruta, "wb") as fh:
                fh.write(html_a_pdf(_cc_html(session, cc)))
            cc.ruta_pdf = ruta
            session.add(cc)
            session.commit()
            session.refresh(cc)
        except Exception:  # noqa: BLE001
            pass
    return cc, True


def filas_cc(session: Session, fac: Facturacion) -> list[dict]:
    """Una fila por CLIENTE con ítems CC en la facturación (NO por ítem): agrupa
    todos sus conceptos del mes en una sola cuenta de cobro (número, total, estado).
    Reutilizado por la sección 'Cuentas de cobro' (módulo cobros)."""
    from .parametros import _cobro
    carpeta = _cobro(session).get("carpeta", "")
    items = session.exec(select(FacturacionItem).where(
        FacturacionItem.facturacion_id == fac.id)).all()
    cc_items = [it for it in items if (it.forma or "").upper() == "CC"]
    # Agrupa los ítems CC por cliente, preservando el orden de aparición de clientes.
    por_cliente: dict = {}
    for it in cc_items:
        por_cliente.setdefault(it.cliente_id, []).append(it)
    # Precarga en bloque (evita N+1): clientes por IN y emisores (catálogo chico) una vez.
    cli_ids = {cid for cid in por_cliente if cid}
    clientes = ({c.id: c for c in session.exec(
                    select(Cliente).options(*cliente_liviano()).where(Cliente.id.in_(cli_ids)))}
                if cli_ids else {})
    emisores = {e.id: e for e in session.exec(select(Emisor))}
    # Cuentas de cobro del mes, indexadas por cliente (una por cliente).
    ccs = {c.cliente_id: c for c in session.exec(
        select(CuentaCobro).where(
            CuentaCobro.anio == fac.anio, CuentaCobro.mes == fac.mes).order_by(CuentaCobro.id)).all()}
    out = []
    for cid, its in por_cliente.items():
        cli = clientes.get(cid) if cid else None
        em = emisores.get(cli.emisor_id) if (cli and cli.emisor_id) else None
        cc = ccs.get(cid)
        conceptos, total = _conceptos_de(its)
        nombre = cli.nombre if cli else (its[0].nombre if its else "")
        concepto_lbl = " · ".join(c["concepto"] for c in conceptos if c["concepto"])
        # N° definitivo (si ya se generó) o el próximo que se asignaría (vista previa).
        proximo = (cc.numero if cc else
                   (f"{em.prefijo}-{em.consecutivo + 1}" if (em and em.prefijo)
                    else (str(em.consecutivo + 1) if em else None)))
        ruta = (cc.ruta_pdf if (cc and cc.ruta_pdf)
                else (_ruta_cc(carpeta, proximo, nombre, fac.anio, fac.mes) if proximo else ""))
        out.append({
            "cliente_id": cid, "cliente": nombre,
            "concepto": concepto_lbl, "conceptos": conceptos, "n_conceptos": len(conceptos),
            "valor": float(cc.valor) if (cc and cc.valor is not None) else float(total),
            "emisor": em.nombre if em else "", "tiene_emisor": bool(em),
            "correo": (cli.email or "").strip() if cli else "",
            "numero": cc.numero if cc else None,
            "proximo_numero": proximo, "ruta": ruta,
            "estado": cc.estado if cc else "pendiente",
            "enviado_at": cc.enviado_at.isoformat() if (cc and cc.enviado_at) else None,
        })
    return out


@router.get("/{fid}/cuentas-cobro")
def cuentas_cobro_estado(fid: int, session: Session = Depends(get_session),
                         _u: Usuario = Depends(require_perm("facturacion.ver"))):
    """Estado de las cuentas de cobro de una facturación (por id)."""
    fac = session.get(Facturacion, fid)
    if not fac:
        raise HTTPException(404, "Facturación no encontrada")
    return filas_cc(session, fac)


@router.post("/{fid}/cliente/{cliente_id}/cuenta-cobro")
def generar_cuenta_cobro(fid: int, cliente_id: int, session: Session = Depends(get_session),
                         u: Usuario = Depends(require_perm("facturacion.editar"))):
    """Genera EN FIRME la cuenta de cobro del cliente (agrupa sus conceptos CC del
    mes; quema el consecutivo)."""
    cc, nueva = _generar_cc(session, fid, cliente_id, u.nombre or u.usuario or "")
    return {"cliente_id": cliente_id, "numero": cc.numero, "estado": cc.estado, "nueva": nueva, "ruta": cc.ruta_pdf}


@router.post("/{fid}/cuentas-cobro/generar")
def generar_cuentas_cobro_lote(fid: int, session: Session = Depends(get_session),
                               u: Usuario = Depends(require_perm("facturacion.editar"))):
    """Genera EN FIRME todas las cuentas de cobro pendientes de la facturación: UNA
    por cliente con ítems CC y emisor asignado que aún no tenga número. Las que ya
    existen o no tienen emisor se omiten; devuelve el conteo y las que no se pudieron."""
    fac = session.get(Facturacion, fid)
    if not fac:
        raise HTTPException(404, "Facturación no encontrada")
    from .parametros import _cobro
    pago = _cobro(session)  # una sola vez para todo el lote
    items = session.exec(select(FacturacionItem).where(
        FacturacionItem.facturacion_id == fid)).all()
    # Clientes con ítems CC (orden de aparición), uno por cliente.
    clientes_cc: dict = {}
    for it in items:
        if (it.forma or "").upper() == "CC" and it.cliente_id:
            clientes_cc.setdefault(it.cliente_id, it.nombre)
    generadas, omitidas, no_generadas = 0, 0, []
    for cid, nombre in clientes_cc.items():
        try:
            _, nueva = _generar_cc(session, fid, cid, u.nombre or u.usuario or "", pago=pago)
            if nueva:
                generadas += 1
            else:
                omitidas += 1
        except HTTPException as e:
            no_generadas.append({"cliente_id": cid, "cliente": nombre, "motivo": e.detail})
    return {"generadas": generadas, "ya_existian": omitidas, "no_generadas": no_generadas}


@router.get("/{fid}/cliente/{cliente_id}/cuenta-cobro.pdf")
def cuenta_cobro_pdf(fid: int, cliente_id: int, session: Session = Depends(get_session),
                     _u: Usuario = Depends(require_perm("facturacion.ver"))):
    """PDF de la cuenta de cobro del cliente (agrupa sus conceptos CC). Si ya se
    generó EN FIRME, usa la foto guardada (número definitivo); si no, es VISTA
    PREVIA con el siguiente consecutivo del emisor (sin consumirlo)."""
    from datetime import date
    from .parametros import _cobro
    from ..cuenta_cobro import html_a_pdf
    from ..models import CuentaCobro

    fac = session.get(Facturacion, fid)
    cc = _cc_existente(session, fac.anio, fac.mes, cliente_id) if fac else None
    if cc:
        html = _cc_html(session, cc)
        numero, nombre_cli = cc.numero, cc.deudor
    else:
        # VISTA PREVIA: arma una CuentaCobro TRANSITORIA (no se persiste) y reusa
        # _cc_html → construir_html vive en un solo lugar (firme y preview iguales).
        items_cc, cli, em, deudor_doc = _cc_validar_cliente(session, fid, cliente_id)
        conceptos, total = _conceptos_de(items_cc)
        numero = f"{em.prefijo}-{em.consecutivo + 1}" if em.prefijo else str(em.consecutivo + 1)
        pago = _cobro(session)
        draft = CuentaCobro(
            numero=numero, fecha=date.today(), deudor=cli.nombre, deudor_doc=deudor_doc,
            emisor_id=em.id, emisor_nombre=em.nombre, emisor_cedula=em.cedula,
            emisor_direccion=em.direccion, emisor_celular=em.celular,
            nota_383=em.nota_383, nota_383_texto=pago["nota_383"], emisor_firma=em.firma,
            conceptos=json.dumps(conceptos, ensure_ascii=False), valor=total,
            pago_banco=pago["banco"], pago_tipo=pago["tipo_cuenta"],
            pago_numero=pago["numero_cuenta"], pago_titular=pago["titular"])
        html = _cc_html(session, draft)
        nombre_cli = cli.nombre
    try:
        pdf = html_a_pdf(html)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"No se pudo generar el PDF: {e}")
    nom = f"CC {numero} - {(nombre_cli or 'cliente')[:40]}.pdf".replace("/", "-")
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{nom}"'})

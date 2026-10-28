"""CRUD de Clientes (con sus hijas: servicios y municipios ICA).

La clave de portal se cifra al guardar y nunca se devuelve (solo un endpoint
de administrador puede revelarla, igual que en Documentos DIAN).
"""
from __future__ import annotations

import io
import os
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from sqlalchemy import func
from sqlmodel import Session, select

from pydantic import BaseModel

from ..db import get_session
from ..deps import get_current_user, require_perm, require_superadmin
from ..models import (Cliente, ClienteClave, ClienteMunicipioIca, ClienteObligacion,
                      ClienteServicio, CuentaCobro, FacturacionItem, ObligacionCumplimiento,
                      ObligacionTipo, Portal, RolServicio, Software, Usuario)
from ..obligaciones_dian import (codigos_resp as _codigos_resp, desincronizado,
                                 diff_mem, dian_tipos, ids_aplicables,
                                 present_map as _present_map)
from ..rut_parser import parse_rut
from ..schemas import (ClienteCreate, ClienteListItem, ClienteRead, ServicioOut,
                       MunicipioIcaOut, ObligacionOut)
from ..security import decrypt, encrypt
from ..util import calcular_dv, componer_nombre

router = APIRouter(prefix="/api/clientes", tags=["clientes"])

_HIJAS = {"password", "servicios", "municipios_ica", "obligaciones"}


def _aplicar_nombre(c: Cliente) -> None:
    """Deriva ``nombre`` según el tipo (Jur = razón social; Nat = componentes)."""
    if c.tipo_contribuyente == "Jur":
        c.tipo_documento = "NT"
        c.nombre = (c.razon_social or c.nombre or "").strip()
    else:
        armado = componer_nombre(c.primer_apellido, c.segundo_apellido,
                                 c.primer_nombre, c.otros_nombres)
        c.nombre = armado or (c.nombre or "").strip()




def _derivar_agente_retenedor(c: Cliente) -> None:
    """Agente de retención sale del RUT (53): responsabilidad 07 o 23.

    No se llena a mano —decisión de Renzo del 13-oct-2026—: antes era un campo
    sin control en la pantalla, con valor solo en los clientes de la carga
    inicial de julio y nulo en los 238 creados después. Lo que llegue en el
    campo desde afuera se descarta.

    Son los mismos dos códigos con que el Liquidador 350 decide a quién ofrece,
    y tienen que seguir siendo los mismos: si discreparan, un cliente podría
    figurar como agente de retención aquí y no poder liquidarse allá.

    Sin responsabilidades cargadas queda en nulo, igual que `_normalizar_regimen_iva`
    no toca el régimen sin ese dato: no se puede afirmar que alguien no es agente
    de retención cuando lo que pasa es que no se sabe.

    En Postgres lo sostiene además un disparador (`db._agente_retenedor_derivado`),
    que cubre toda vía de escritura. Esto lo repite aquí para que valga también
    en SQLite y para que quede probado.
    """
    cods = _codigos_resp(c)
    c.agente_retenedor = bool({"07", "23"} & set(cods)) if cods else None


def _normalizar_regimen_iva(c: Cliente) -> None:
    """Valida y ajusta Régimen y Periodicidad IVA según las responsabilidades (RUT 53):
      · 47 (RST) => Régimen "Simple" + IVA "Anual".
      · sin 47   => Régimen "Ordinario"; sin 48 no hay IVA; con 48 no puede ser "Anual".
    Si NO hay responsabilidades cargadas (RUT no importado) no se toca nada: no se
    puede inferir el régimen sin ese dato."""
    cods = _codigos_resp(c)
    if not cods:
        return
    rst = "47" in cods
    c.regimen = "Simple" if rst else "Ordinario"
    if rst:
        c.periodicidad_iva = "Anual"
    elif "48" not in cods:
        c.periodicidad_iva = ""
    elif c.periodicidad_iva == "Anual":     # 48 sin 47: la elige el usuario (Bim/Cuat), no Anual
        c.periodicidad_iva = ""


def _sync_hijas(session: Session, cliente_id: int, data: ClienteCreate,
                sync_servicios: bool = True) -> None:
    """Reemplaza municipios ICA (y servicios) del cliente por los del payload.
    ``sync_servicios=False`` deja intactos los servicios (dato comercial que el
    Administrador no puede modificar)."""
    # Al reemplazar los servicios preservamos la marca de facturado: se guarda
    # por (rol, período) para que un guardado de la ficha no borre el estado que
    # fija el flujo de facturación (facturacion_id es autoridad del servidor).
    marcas_fact: dict[tuple, int] = {}
    if sync_servicios:
        for viejo in session.exec(
                select(ClienteServicio).where(ClienteServicio.cliente_id == cliente_id)).all():
            if viejo.facturacion_id:
                marcas_fact[(viejo.rol_servicio_id, viejo.periodo)] = viejo.facturacion_id
            session.delete(viejo)
    for viejo in session.exec(
            select(ClienteMunicipioIca).where(ClienteMunicipioIca.cliente_id == cliente_id)).all():
        session.delete(viejo)
    # No permitir QUITAR una obligación que ya tiene entradas en el Calendario (cumplimientos).
    _exist = {co.obligacion_tipo_id for co in session.exec(
        select(ClienteObligacion).where(ClienteObligacion.cliente_id == cliente_id)).all()}
    _entra = {ob.obligacion_tipo_id for ob in data.obligaciones}
    _bloq = []
    for _otid in (_exist - _entra):
        if _otid is None:
            continue
        _n = len(session.exec(select(ObligacionCumplimiento).where(
            ObligacionCumplimiento.cliente_id == cliente_id,
            ObligacionCumplimiento.obligacion_tipo_id == _otid)).all())
        if _n:
            _t = session.get(ObligacionTipo, _otid)
            _bloq.append(f"{_t.descripcion if _t else _otid} ({_n})")
    if _bloq:
        raise HTTPException(409, "No se puede quitar: tiene registros de cumplimiento/revisión "
                                 "en el Calendario — " + "; ".join(_bloq)
                                 + ". Desactívala o elimina esos registros primero.")
    for viejo in session.exec(
            select(ClienteObligacion).where(ClienteObligacion.cliente_id == cliente_id)).all():
        session.delete(viejo)
    if sync_servicios:
        for s in data.servicios:
            fid = marcas_fact.get((s.rol_servicio_id, s.periodo))
            session.add(ClienteServicio(cliente_id=cliente_id, facturacion_id=fid,
                                        **s.model_dump()))
    for m in data.municipios_ica:
        session.add(ClienteMunicipioIca(cliente_id=cliente_id, **m.model_dump()))
    # La periodicidad de la obligación del cliente NO se edita por cliente (no hay
    # UI para ella): es un atributo del catálogo. Fuente única = ObligacionTipo, para
    # que el Calendario genere los vencimientos con la frecuencia correcta. Ignoramos
    # lo que mande el front (histórico: caía "Mensual" por el default del schema).
    if data.obligaciones:
        _tids = {ob.obligacion_tipo_id for ob in data.obligaciones if ob.obligacion_tipo_id}
        _per = {t.id: t.periodicidad for t in session.exec(
            select(ObligacionTipo).where(ObligacionTipo.id.in_(_tids))).all()} if _tids else {}
        for ob in data.obligaciones:
            d = ob.model_dump()
            d["periodicidad"] = (_per.get(ob.obligacion_tipo_id) or "").strip() or "Anual"
            session.add(ClienteObligacion(cliente_id=cliente_id, **d))


# Campos comerciales que el Administrador NO puede ver ni modificar (solo superusuario).
# Campos de la pestaña "Comercial" (solo el Super Usuario los edita; para los demás
# se ocultan al leer y se revierten al guardar). SOLO los de esa pestaña: p. ej.
# `ruta_trabajo` vive en la pestaña Gestión, así que NO va aquí (cualquier rol con
# `clientes.editar` debe poder guardarla). Los servicios (hija comercial) se controlan
# aparte con `sync_servicios`.
_COMERCIAL = ("cliente_desde", "estado", "archivo_propuesta", "emisor_id")

# Rutas de archivo/carpeta por cliente: se GUARDAN portables (token %OneDrive%)
# para que sirvan en CUALQUIER equipo, y se EXPANDEN al OneDrive de ESTE equipo al
# leerlas (para mostrarlas y abrir el explorador ahí). Una ruta absoluta de otro
# equipo (p. ej. C:\Users\COMPUMAX\OneDrive…\X) no existe en los demás y rompía el
# selector: por eso nunca se guarda absoluta.
_PATH_FIELDS = ("ruta_trabajo", "ruta_generador", "archivo_propuesta")


def _ruta_portable(p: str) -> str:
    """Absoluta bajo OneDrive → `%OneDriveCommercial%\\…` (o `%OneDrive%`); si ya es
    token o no cuelga de OneDrive, la deja igual. Idempotente."""
    ruta = (p or "").strip()
    if not ruta or ruta.startswith("%"):
        return ruta
    for var in ("OneDriveCommercial", "OneDrive"):
        base = (os.environ.get(var) or "").rstrip("\\/")
        if base and ruta.lower().startswith(base.lower()):
            return f"%{var}%" + ruta[len(base):]
    return ruta


def _ruta_expandir(p: str) -> str:
    """Token `%OneDrive%`/`~` → ruta absoluta de ESTE equipo (mostrar/abrir). Si el
    token usa una variable de OneDrive que este equipo no tiene pero sí la otra, la
    cambia (comercial y personal apuntan al mismo OneDrive de la firma)."""
    ruta = (p or "").strip()
    if not ruta:
        return ruta
    if ruta.startswith("%OneDriveCommercial%") and not os.environ.get("OneDriveCommercial") and os.environ.get("OneDrive"):
        ruta = "%OneDrive%" + ruta[len("%OneDriveCommercial%"):]
    elif ruta.startswith("%OneDrive%") and not os.environ.get("OneDrive") and os.environ.get("OneDriveCommercial"):
        ruta = "%OneDriveCommercial%" + ruta[len("%OneDrive%"):]
    return os.path.expanduser(os.path.expandvars(ruta))


def _portabilizar_rutas(c: Cliente) -> None:
    """Deja las rutas del cliente en forma PORTABLE antes de persistir."""
    for k in _PATH_FIELDS:
        setattr(c, k, _ruta_portable(getattr(c, k, "") or ""))


def _leer(session: Session, c: Cliente, ver_comercial: bool = True) -> ClienteRead:
    servicios = session.exec(
        select(ClienteServicio).where(ClienteServicio.cliente_id == c.id)).all()
    municipios = session.exec(
        select(ClienteMunicipioIca).where(ClienteMunicipioIca.cliente_id == c.id)).all()
    obligaciones = session.exec(
        select(ClienteObligacion).where(ClienteObligacion.cliente_id == c.id)).all()
    base = c.model_dump()
    base.pop("password_enc", None)
    if not ver_comercial:                       # oculta lo comercial a los Administradores
        for k in _COMERCIAL:                     # vacío por tipo: "" para texto, None para numéricos (p.ej. emisor_id)
            base[k] = "" if isinstance(base.get(k), str) else None
        servicios = []
    for k in _PATH_FIELDS:                        # token portable → ruta de ESTE equipo (display + explorador)
        if base.get(k):
            base[k] = _ruta_expandir(base[k])
    return ClienteRead(
        **base,
        tiene_password=bool(c.password_enc),
        servicios=[ServicioOut.model_validate(s, from_attributes=True) for s in servicios],
        municipios_ica=[MunicipioIcaOut.model_validate(m, from_attributes=True) for m in municipios],
        obligaciones=[ObligacionOut.model_validate(o, from_attributes=True) for o in obligaciones],
    )


@router.get("", response_model=list[ClienteListItem])
def listar(session: Session = Depends(get_session), user: Usuario = Depends(get_current_user)):
    # Alcance por asignación: cada usuario ve SOLO sus clientes asignados
    # (responsable/coordinador de impuestos o financiero). Admin/Super ven todos.
    ids, ocultos = _alcance_de(session, user)    # asignación + confidenciales, una pasada
    # Sin columnas pesadas (logo/objeto_social/notas): el listado no las usa.
    q = select(Cliente).options(*cliente_liviano())
    if ids is not None:
        if not ids:
            return []
        q = q.where(Cliente.id.in_(ids))
    if ocultos:
        q = q.where(Cliente.id.notin_(ocultos))
    clientes = session.exec(q).all()
    _usuarios = session.exec(select(Usuario)).all()
    responsables = {u.id: u.nombre for u in _usuarios}
    apodos = {u.id: (u.apodo or "") for u in _usuarios}
    softwares = {s.id: s.descripcion for s in session.exec(select(Software)).all()}
    # Conteos de hijas agregados en la BD (no se traen todas las filas).
    n_srv: dict[int, int] = dict(session.exec(
        select(ClienteServicio.cliente_id, func.count()).group_by(ClienteServicio.cliente_id)).all())
    n_ica: dict[int, int] = dict(session.exec(
        select(ClienteMunicipioIca.cliente_id, func.count()).group_by(ClienteMunicipioIca.cliente_id)).all())
    # Diferencias DIAN por revisar (chip "por revisar"): precargado en memoria
    # (tipos DIAN + obligaciones presentes por cliente) para no consultar por fila.
    dian = session.exec(select(ObligacionTipo).where(ObligacionTipo.origen == "dian")).all()
    present_map: dict[int, set[int]] = {}
    for cid, otid in session.exec(
            select(ClienteObligacion.cliente_id, ClienteObligacion.obligacion_tipo_id)).all():
        present_map.setdefault(cid, set()).add(otid)
    items = []
    for c in clientes:
        items.append(ClienteListItem(
            obligaciones_por_revisar=desincronizado(
                c, _codigos_resp(c), dian, present_map.get(c.id, set())),
            id=c.id, num_id=c.num_id, dv=calcular_dv(c.num_id), nombre=c.nombre,
            email=c.email,
            tipo_contribuyente=c.tipo_contribuyente, tipo_documento=c.tipo_documento,
            regimen=c.regimen, estado=c.estado, responsable_impuestos_id=c.responsable_impuestos_id,
            responsable_nombre=responsables.get(c.responsable_impuestos_id, ""),
            software_desc=softwares.get(c.software_id, ""),
            n_servicios=n_srv.get(c.id, 0), n_municipios=n_ica.get(c.id, 0),
            tiene_password=bool(c.password_enc), rut_actualizado=c.rut_actualizado or "",
            confidencial=c.confidencial,
            resp_imp_apodo=apodos.get(c.responsable_impuestos_id, ""),
            coord_imp_apodo=apodos.get(c.coordinador_impuestos_id, ""),
            resp_fin_apodo=apodos.get(c.responsable_financiero_id, ""),
            coord_fin_apodo=apodos.get(c.coordinador_financiero_id, ""),
            resp_imp_nombre=responsables.get(c.responsable_impuestos_id, ""),
            coord_imp_nombre=responsables.get(c.coordinador_impuestos_id, ""),
            resp_fin_nombre=responsables.get(c.responsable_financiero_id, ""),
            coord_fin_nombre=responsables.get(c.coordinador_financiero_id, "")))
    return items


@router.post("/rut/parse")
async def rut_parse(file: UploadFile = File(...), _u=Depends(get_current_user)):
    """Lee un RUT (PDF de la DIAN) y devuelve las casillas relevantes. No guarda:
    el frontend arma el diff contra la ficha y el usuario decide qué aplicar."""
    data = await file.read()
    if not data:
        raise HTTPException(400, "El archivo está vacío")
    try:
        leido = parse_rut(data)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el RUT: {exc}")
    if not leido:
        raise HTTPException(422, "No se reconoció el formato del RUT (¿PDF escaneado?)")
    return leido


@router.get("/rut/parse-path")
def rut_parse_path(path: str, _u=Depends(get_current_user)):
    """Igual que /rut/parse pero leyendo el PDF desde una ruta del disco local
    (la que devuelve el selector nativo). El backend corre en la máquina del
    usuario, así que puede leer el archivo elegido."""
    p = Path(path)
    if not path or not p.is_file():
        raise HTTPException(400, "Ruta de archivo no válida")
    if p.suffix.lower() != ".pdf":
        raise HTTPException(400, "El archivo no es un PDF")
    try:
        data = p.read_bytes()
    except OSError as exc:
        raise HTTPException(400, f"No se pudo leer el archivo: {exc}")
    if not data:
        raise HTTPException(400, "El archivo está vacío")
    try:
        leido = parse_rut(data)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el RUT: {exc}")
    if not leido:
        raise HTTPException(422, "No se reconoció el formato del RUT (¿PDF escaneado?)")
    return leido


@router.get("/{cid}", response_model=ClienteRead)
def obtener(cid: int, session: Session = Depends(get_session), u: Usuario = Depends(get_current_user)):
    c = session.get(Cliente, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    ids, ocultos = _alcance_de(session, u)   # alcance por asignación (None=todos) + confidenciales
    if cid in ocultos:   # confidencial que no puede ver (ni asignado ni supervisor)
        raise HTTPException(404, "Cliente no encontrado")
    if ids is not None and cid not in ids:
        raise HTTPException(403, "Este cliente no está asignado a tu usuario")
    return _leer(session, c, ver_comercial=(u.rol == "superadmin"))


@router.get("/{cid}/obligacion/{otid}/entradas-calendario")
def entradas_calendario(cid: int, otid: int, session: Session = Depends(get_session),
                        _u: Usuario = Depends(get_current_user)):
    """¿El cliente ya tiene entradas (cumplimientos) en el Calendario para esta
    obligación? Si las tiene, no se puede quitarla (se perderían esos registros)."""
    n = len(session.exec(select(ObligacionCumplimiento).where(
        ObligacionCumplimiento.cliente_id == cid,
        ObligacionCumplimiento.obligacion_tipo_id == otid)).all())
    return {"tiene": n > 0, "n": n}


def _descifrar(enc: str) -> str:
    """Descifra una clave guardada; responde 422 (no 500) si la llave Fernet no coincide.
    Helper PRIVADO (sin ruta): nunca debe exponerse como endpoint — sería un oráculo
    de descifrado público. El endpoint autenticado es `ver_password`."""
    try:
        return decrypt(enc)
    except Exception:  # noqa: BLE001  (InvalidToken u otros: no reventar con 500)
        raise HTTPException(422, "No se pudo descifrar la clave (llave de cifrado no coincide).")


@router.get("/{cid}/password")
def ver_password(cid: int, session: Session = Depends(get_session), _admin=Depends(require_perm("claves.revelar"))):
    c = session.get(Cliente, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    if not c.password_enc:
        return {"password": ""}
    return {"password": _descifrar(c.password_enc)}


# ============================================================ Claves de portales
class ClaveIn(BaseModel):
    portal_id: int | None = None
    comentario: str = ""
    url: str = ""
    odbc_nombre: str = ""
    odbc_dbf: str = ""
    odbc_ip: str = ""
    usuario: str = ""
    clave: str = ""     # en PUT: vacío = no cambiar la clave existente


def _clave_base(k: ClienteClave, portales: dict) -> dict:
    """Campos comunes de una clave (los que comparten la vista por cliente y la
    global). Cada vista extiende este dict con lo suyo."""
    return {"id": k.id, "portal": portales.get(k.portal_id, ""),
            "comentario": k.comentario, "url": k.url, "usuario": k.usuario,
            "tiene_clave": bool(k.clave_enc),
            "actualizado_at": k.actualizado_at.isoformat() if k.actualizado_at else None}


def _clave_dict(k: ClienteClave, portales: dict) -> dict:
    """Vista por cliente (ficha + modal): núcleo + portal_id + datos ODBC."""
    return {**_clave_base(k, portales), "portal_id": k.portal_id,
            "odbc_nombre": k.odbc_nombre, "odbc_dbf": k.odbc_dbf, "odbc_ip": k.odbc_ip}


# Reglas de alcance/confidencialidad: fuente única en ``scope`` (ver scope.py).
from ..scope import (alcance_de as _alcance_de, cliente_liviano,  # noqa: E402
                     ve_confidencial as _ve_confidencial)


@router.get("/claves/todas")
def listar_todas_claves(session: Session = Depends(get_session),
                        user: Usuario = Depends(require_perm("claves.ver"))):
    """Todas las claves visibles para el usuario (SIN la clave en claro).
    Admin/Super ven todas; los demás solo las de sus clientes asignados.
    Alimenta la vista global 'Claves'. Se revela una a una con /claves/{kid}/revelar."""
    permitidos, ocultos = _alcance_de(session, user)
    portales = {p.id: p.descripcion for p in session.exec(select(Portal)).all()}
    clientes = {c.id: c for c in session.exec(select(Cliente).options(*cliente_liviano())).all()}
    filas = []
    for k in session.exec(select(ClienteClave)).all():
        if permitidos is not None and k.cliente_id not in permitidos:
            continue
        if k.cliente_id in ocultos:                # claves de confidenciales no visibles
            continue
        c = clientes.get(k.cliente_id)
        filas.append({
            **_clave_base(k, portales), "cliente_id": k.cliente_id,
            "num_id": c.num_id if c else "", "cliente": c.nombre if c else "(cliente eliminado)",
            "estado": c.estado if c else "", "tipo": c.tipo_contribuyente if c else "",
            "rl": c.representante_nombre if c else "",
        })
    filas.sort(key=lambda r: ((r["cliente"] or "").lower(), (r["portal"] or "").lower(),
                              (r["usuario"] or "").lower()))
    return filas


@router.get("/claves/{kid}/revelar")
def revelar_clave_global(kid: int, session: Session = Depends(get_session),
                         user: Usuario = Depends(require_perm("claves.revelar"))):
    """Revela una clave por su id (vista global). Admin/Super, o el usuario
    asignado al cliente de esa clave."""
    k = session.get(ClienteClave, kid)
    if not k:
        raise HTTPException(404, "Clave no encontrada")
    permitidos, ocultos = _alcance_de(session, user)
    if k.cliente_id in ocultos:   # confidencial no visible para este usuario
        raise HTTPException(404, "Clave no encontrada")
    if permitidos is not None and k.cliente_id not in permitidos:
        raise HTTPException(403, "No tiene acceso a las claves de este cliente")
    if not k.clave_enc:
        return {"clave": ""}
    return {"clave": _descifrar(k.clave_enc)}


@router.get("/claves/reveladas")
def revelar_todas_claves(session: Session = Depends(get_session),
                         user: Usuario = Depends(require_perm("claves.revelar"))):
    """Devuelve {id, clave} en claro de todas las claves visibles para el usuario
    (Admin/Super o sus clientes asignados). Alimenta el 'mostrar toda la columna'."""
    permitidos, ocultos = _alcance_de(session, user)
    out = []
    for k in session.exec(select(ClienteClave)).all():
        if permitidos is not None and k.cliente_id not in permitidos:
            continue
        if k.cliente_id in ocultos:                # claves de confidenciales no visibles
            continue
        if not k.clave_enc:
            continue
        try:
            out.append({"id": k.id, "clave": decrypt(k.clave_enc)})
        except Exception:  # noqa: BLE001
            out.append({"id": k.id, "clave": ""})
    return out


@router.get("/claves/exportar")
def exportar_claves(session: Session = Depends(get_session),
                    user: Usuario = Depends(require_perm("claves.exportar"))):
    """Exporta a Excel las claves visibles para el usuario (con la clave en claro).
    Respeta el filtro por asignación (cada quien exporta solo lo suyo)."""
    from ..export_xlsx import tabla_simple

    permitidos, ocultos = _alcance_de(session, user)
    portales = {p.id: p.descripcion for p in session.exec(select(Portal)).all()}
    clientes = {c.id: c for c in session.exec(select(Cliente).options(*cliente_liviano())).all()}
    filas = []
    for k in session.exec(select(ClienteClave)).all():
        if permitidos is not None and k.cliente_id not in permitidos:
            continue
        if k.cliente_id in ocultos:                # confidenciales no visibles
            continue
        c = clientes.get(k.cliente_id)
        clave = ""
        if k.clave_enc:
            try:
                clave = decrypt(k.clave_enc)
            except Exception:  # noqa: BLE001
                clave = "(no descifrable)"
        filas.append((c.num_id if c else "", c.nombre if c else "(cliente eliminado)",
                      portales.get(k.portal_id, ""), k.url, k.usuario, k.comentario, clave,
                      k.actualizado_at.strftime("%Y-%m-%d") if k.actualizado_at else ""))
    filas.sort(key=lambda f: ((f[1] or "").lower(), (f[2] or "").lower(), (f[3] or "").lower()))

    encab = ["NIT", "Cliente", "Portal", "URL", "Usuario", "Comentario", "Clave", "Actualizado"]
    return tabla_simple("Claves_clientes.xlsx", "Claves", encab, filas, [14, 40, 22, 30, 24, 28, 22, 13])


@router.delete("/claves/{kid}", status_code=204)
def borrar_clave_global(kid: int, session: Session = Depends(get_session),
                        _u: Usuario = Depends(require_perm("claves.eliminar"))):
    """Elimina una clave por su id desde la vista global. Solo Super Usuario."""
    k = session.get(ClienteClave, kid)
    if not k:
        raise HTTPException(404, "Clave no encontrada")
    session.delete(k)
    session.commit()


@router.get("/{cid}/claves")
def listar_claves(cid: int, session: Session = Depends(get_session), _admin=Depends(require_perm("claves.ver"))):
    """Claves del cliente (SIN la clave en claro: solo si la tiene). Admin/Super."""
    portales = {p.id: p.descripcion for p in session.exec(select(Portal)).all()}
    ks = session.exec(select(ClienteClave).where(ClienteClave.cliente_id == cid)).all()
    ks.sort(key=lambda k: (portales.get(k.portal_id, "").lower(), k.usuario.lower()))
    return [_clave_dict(k, portales) for k in ks]


@router.post("/{cid}/claves", status_code=201)
def crear_clave(cid: int, data: ClaveIn, session: Session = Depends(get_session),
                _admin=Depends(require_perm("claves.crear"))):
    if not session.get(Cliente, cid):
        raise HTTPException(404, "Cliente no encontrado")
    k = ClienteClave(cliente_id=cid, portal_id=data.portal_id, comentario=data.comentario.strip(),
                     url=data.url.strip(), odbc_nombre=data.odbc_nombre.strip(),
                     odbc_dbf=data.odbc_dbf.strip(), odbc_ip=data.odbc_ip.strip(),
                     usuario=data.usuario.strip(),
                     clave_enc=encrypt(data.clave) if data.clave else "")
    session.add(k)
    session.commit()
    session.refresh(k)
    portales = {p.id: p.descripcion for p in session.exec(select(Portal)).all()}
    return _clave_dict(k, portales)


@router.put("/{cid}/claves/{kid}")
def actualizar_clave(cid: int, kid: int, data: ClaveIn, session: Session = Depends(get_session),
                     _admin=Depends(require_perm("claves.editar"))):
    k = session.get(ClienteClave, kid)
    if not k or k.cliente_id != cid:
        raise HTTPException(404, "Clave no encontrada")
    k.portal_id = data.portal_id
    k.comentario = data.comentario.strip()
    k.url = data.url.strip()
    k.odbc_nombre = data.odbc_nombre.strip()
    k.odbc_dbf = data.odbc_dbf.strip()
    k.odbc_ip = data.odbc_ip.strip()
    k.usuario = data.usuario.strip()
    if data.clave:                       # vacío = conservar la clave actual
        k.clave_enc = encrypt(data.clave)
    k.actualizado_at = datetime.utcnow()
    session.add(k)
    session.commit()
    session.refresh(k)
    portales = {p.id: p.descripcion for p in session.exec(select(Portal)).all()}
    return _clave_dict(k, portales)


@router.delete("/{cid}/claves/{kid}", status_code=204)
def borrar_clave(cid: int, kid: int, session: Session = Depends(get_session),
                 _admin=Depends(require_perm("claves.eliminar"))):
    k = session.get(ClienteClave, kid)
    if not k or k.cliente_id != cid:
        raise HTTPException(404, "Clave no encontrada")
    session.delete(k)
    session.commit()


@router.get("/{cid}/claves/{kid}/revelar")
def revelar_clave(cid: int, kid: int, session: Session = Depends(get_session),
                  _admin=Depends(require_perm("claves.revelar"))):
    """Devuelve la clave en claro (Admin/Super), igual que la clave DIAN."""
    k = session.get(ClienteClave, kid)
    if not k or k.cliente_id != cid:
        raise HTTPException(404, "Clave no encontrada")
    if not k.clave_enc:
        return {"clave": ""}
    return {"clave": _descifrar(k.clave_enc)}


@router.post("", response_model=ClienteRead, status_code=201)
def crear(data: ClienteCreate, session: Session = Depends(get_session),
          user: Usuario = Depends(require_perm("clientes.crear"))):
    if session.exec(select(Cliente).where(Cliente.num_id == data.num_id)).first():
        raise HTTPException(409, "Ya existe un cliente con ese número de documento")
    es_super = user.rol == "superadmin"
    c = Cliente(**data.model_dump(exclude=_HIJAS))
    if not _ve_confidencial(session, user):     # solo supervisores marcan confidencial
        c.confidencial = False
    if not es_super:                            # el Administrador no fija datos comerciales
        for k in _COMERCIAL:
            setattr(c, k, "")
    _aplicar_nombre(c)
    _normalizar_regimen_iva(c)
    _derivar_agente_retenedor(c)
    _portabilizar_rutas(c)
    c.dv = calcular_dv(c.num_id)
    if data.password:
        c.password_enc = encrypt(data.password)
    c.creado_por = user.usuario or user.nombre
    session.add(c)
    session.commit()
    session.refresh(c)
    _sync_hijas(session, c.id, data, sync_servicios=es_super)
    session.commit()
    # Las obligaciones DIAN NO se derivan al guardar: se revisan y aplican a mano
    # (botón "Revisar obligaciones"), nunca en silencio.
    return _leer(session, c, ver_comercial=es_super)


@router.put("/{cid}", response_model=ClienteRead)
def actualizar(cid: int, data: ClienteCreate, session: Session = Depends(get_session),
               user: Usuario = Depends(require_perm("clientes.editar"))):
    c = session.get(Cliente, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    choque = session.exec(select(Cliente).where(Cliente.num_id == data.num_id)).first()
    if choque and choque.id != cid:
        raise HTTPException(409, "Ya existe otro cliente con ese número de documento")
    es_super = user.rol == "superadmin"
    ve_conf = _ve_confidencial(session, user)
    com_prev = {k: getattr(c, k) for k in _COMERCIAL} if not es_super else None
    conf_prev = c.confidencial
    for k, v in data.model_dump(exclude=_HIJAS).items():
        setattr(c, k, v)
    if not ve_conf:                             # solo supervisores cambian el flag confidencial
        c.confidencial = conf_prev
    if not es_super:                            # revierte lo comercial: el Administrador no lo modifica
        for k, v in com_prev.items():
            setattr(c, k, v)
    _aplicar_nombre(c)
    _normalizar_regimen_iva(c)
    _derivar_agente_retenedor(c)
    _portabilizar_rutas(c)
    c.dv = calcular_dv(c.num_id)
    if data.password:                       # solo si viene una clave nueva
        c.password_enc = encrypt(data.password)
    c.actualizado_at = datetime.utcnow()
    c.actualizado_por = user.usuario or user.nombre
    session.add(c)
    # Validar+sincronizar las hijas ANTES de commitear: si `_sync_hijas` lanza 409
    # (p.ej. quitar una obligación que ya tiene cumplimientos), no debe quedar un
    # guardado PARCIAL de los campos escalares. Una sola transacción, un solo commit.
    _sync_hijas(session, c.id, data, sync_servicios=es_super)
    session.commit()
    # Las obligaciones DIAN NO se derivan al guardar: se revisan y aplican a mano.
    session.refresh(c)
    return _leer(session, c, ver_comercial=es_super)


@router.post("/derivar-preview")
def derivar_preview(data: ClienteCreate, session: Session = Depends(get_session),
                    _u: Usuario = Depends(get_current_user)):
    """Vista previa para el modal 'Revisar obligaciones' de la ficha: IDs de las
    obligaciones DIAN que aplican al borrador (posiblemente sin guardar), SIN
    persistir nada. El frontend calcula el diff (faltan/sobran) contra lo que la
    ficha tiene. Fuente única = la misma regla del backend, sin drift."""
    c = Cliente(**data.model_dump(exclude=_HIJAS))
    return {"aplican": ids_aplicables(session, c)}


# --- Reconciliación DIAN en lote (Catálogos → Obligaciones), REVISABLE ---------
def _diff_lote(session: Session):
    """Resumen del diff DIAN de TODOS los clientes (para el preview del lote).
    Precarga catálogo DIAN y presentes de todos una sola vez (sin N+1)."""
    clientes = session.exec(select(Cliente).options(*cliente_liviano())).all()
    dian = dian_tipos(session)
    presentes = _present_map(session)
    con_faltan = faltan_total = con_sobran = sobran_total = sin_rut = 0
    for c in clientes:
        resp = _codigos_resp(c)
        if not resp:
            sin_rut += 1
            continue
        faltan, sobran = diff_mem(c, resp, dian, presentes.get(c.id, set()))
        if faltan:
            con_faltan += 1
            faltan_total += len(faltan)
        if sobran:
            con_sobran += 1
            sobran_total += len(sobran)
    return {"clientes": len(clientes), "sin_rut": sin_rut,
            "con_faltan": con_faltan, "faltan_total": faltan_total,
            "con_sobran": con_sobran, "sobran_total": sobran_total}


@router.get("/dian-diff-todos")
def dian_diff_todos(session: Session = Depends(get_session),
                    _su: Usuario = Depends(require_superadmin)):
    """Preview del lote: cuántos clientes tienen faltantes y cuántos sobrantes.
    Solo Super Usuario (acción global sobre todos los clientes)."""
    return _diff_lote(session)


@router.post("/dian-agregar-faltantes-todos")
def dian_agregar_faltantes_todos(session: Session = Depends(get_session),
                                 _su: Usuario = Depends(require_superadmin)):
    """Agrega a TODOS los clientes con RUT las obligaciones DIAN faltantes (nunca
    quita). Conserva las manuales. Solo Super Usuario."""
    clientes = session.exec(select(Cliente).options(*cliente_liviano())).all()
    dian = dian_tipos(session)
    by_id = {t.id: t for t in dian}
    presentes = _present_map(session)
    tocados = agregadas = 0
    for c in clientes:
        resp = _codigos_resp(c)
        if not resp:
            continue
        faltan, _ = diff_mem(c, resp, dian, presentes.get(c.id, set()))
        if not faltan:
            continue
        for tid in faltan:
            t = by_id.get(tid)
            session.add(ClienteObligacion(
                cliente_id=c.id, obligacion_tipo_id=tid,
                periodicidad=(t.periodicidad if t else "") or "Anual", desde="", observaciones=""))
        tocados += 1
        agregadas += len(faltan)
    session.commit()
    return {"clientes_tocados": tocados, "agregadas": agregadas}


@router.post("/dian-quitar-sobrantes-todos")
def dian_quitar_sobrantes_todos(session: Session = Depends(get_session),
                                _su: Usuario = Depends(require_superadmin)):
    """Quita de TODOS los clientes las obligaciones DIAN que ya no aplican (nunca
    toca manuales). Acción explícita con confirmación. Solo Super Usuario."""
    clientes = session.exec(select(Cliente).options(*cliente_liviano())).all()
    dian = dian_tipos(session)
    # Filas presentes agrupadas por cliente, una sola consulta (sirve de present_map y para borrar).
    obs_by_client: dict[int, list[ClienteObligacion]] = {}
    for ob in session.exec(select(ClienteObligacion)).all():
        obs_by_client.setdefault(ob.cliente_id, []).append(ob)
    tocados = quitadas = 0
    for c in clientes:
        resp = _codigos_resp(c)
        if not resp:
            continue
        filas = obs_by_client.get(c.id, [])
        present_ids = {ob.obligacion_tipo_id for ob in filas}
        _, sobran = diff_mem(c, resp, dian, present_ids)
        if not sobran:
            continue
        s = set(sobran)
        n = 0
        for ob in filas:
            if ob.obligacion_tipo_id in s:
                session.delete(ob)
                n += 1
        if n:
            tocados += 1
            quitadas += n
    session.commit()
    return {"clientes_tocados": tocados, "quitadas": quitadas}


@router.delete("/{cid}", status_code=204)
def borrar(cid: int, session: Session = Depends(get_session), _admin=Depends(require_perm("clientes.eliminar"))):
    c = session.get(Cliente, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    # No se elimina un cliente con HISTORIA FINANCIERA: las cuentas de cobro y los
    # ítems de facturación son registros inmutables (consecutivo quemado). Se bloquea
    # con mensaje claro, como `responsables.borrar` cuando hay referencias.
    n_cc = len(session.exec(select(CuentaCobro.id).where(CuentaCobro.cliente_id == cid)).all())
    n_fi = len(session.exec(select(FacturacionItem.id).where(FacturacionItem.cliente_id == cid)).all())
    if n_cc or n_fi:
        partes = []
        if n_cc: partes.append(f"{n_cc} cuenta(s) de cobro")
        if n_fi: partes.append(f"{n_fi} ítem(s) de facturación")
        raise HTTPException(409, "No se puede eliminar: el cliente tiene "
                            + " y ".join(partes) + ". Desactívalo (estado Inactivo) en su lugar.")
    # Hijas de configuración: se borran en la MISMA transacción (antes solo se
    # borraban Servicios e ICA → las demás FK reventaban con 500 en Postgres).
    for modelo in (ClienteServicio, ClienteMunicipioIca, ClienteObligacion,
                   ClienteClave, ObligacionCumplimiento):
        for row in session.exec(select(modelo).where(modelo.cliente_id == cid)).all():
            session.delete(row)
    session.delete(c)
    session.commit()


# ============================================================================
#  Importador de facturación (cobros programados) desde Excel
#  Crea filas ClienteServicio cruzando por NIT. Solo Super Usuarios.
# ============================================================================
_solo_dig = lambda s: re.sub(r"\D", "", str(s or ""))


def _norm_hdr(s: str) -> str:
    t = str(s or "").strip().lower()
    t = (t.replace("á", "a").replace("é", "e").replace("í", "i")
          .replace("ó", "o").replace("ú", "u").replace("%", " "))
    return re.sub(r"[^a-z0-9]+", "", t)


# sinónimos de encabezado -> clave canónica
_COLS = {
    "nit": {"nit", "numid", "numeroid", "identificacion", "documento", "cedula"},
    "nombre": {"nombre", "tercero", "cliente", "razonsocial"},
    "rol": {"rol", "rolservicio", "roldelservicio", "servicio", "concepto"},
    "forma": {"forma", "formafacturacion", "tipo", "facturacc", "factura"},
    "periodo": {"periodo", "mes", "periodofacturacion"},
    "valor": {"valor", "valorservicio", "vlr", "monto", "precio"},
    "vendedor": {"vendedor", "comercial", "asesor", "ejecutivo"},
    "pct": {"comision", "pctcomision", "porcentajecomision", "pcomision", "comisionpct"},
    "meses": {"meses", "cantidad", "nmeses", "cuotas", "periodos", "repeticiones"},
}


def _detectar_columnas(fila_hdr: list) -> dict:
    """Mapea nombre-de-columna canónico -> índice en la hoja."""
    mapa: dict[str, int] = {}
    for i, celda in enumerate(fila_hdr):
        h = _norm_hdr(celda)
        if not h:
            continue
        for clave, alias in _COLS.items():
            if clave in mapa:
                continue
            if h in alias:
                mapa[clave] = i
                break
    return mapa


def _norm_forma(v) -> str:
    t = str(v or "").strip().lower()
    if t.startswith("cc") or "cuenta" in t or t == "c":
        return "CC"
    return "Fa"


def _norm_periodo(v) -> str:
    if isinstance(v, datetime):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip()
    if s.endswith(".0"):          # 202608.0 (número leído como float)
        s = s[:-2]
    anio = mes = 0
    m = re.match(r"^(\d{4})[-/.](\d{1,2})$", s)          # 2026-07
    if m:
        anio, mes = int(m.group(1)), int(m.group(2))
    else:
        m = re.match(r"^(\d{1,2})[-/.](\d{4})$", s)      # 07-2026
        if m:
            anio, mes = int(m.group(2)), int(m.group(1))
        else:
            m = re.match(r"^(\d{4})(\d{2})$", s)         # 202608 (AAAAMM)
            if m:
                anio, mes = int(m.group(1)), int(m.group(2))
    if 1 <= mes <= 12 and 2000 <= anio <= 2100:
        return f"{anio:04d}-{mes:02d}"
    return ""


def _money(v) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = re.sub(r"[^\d,.\-]", "", str(v).strip())
    if not s:
        return None
    if "," in s and "." in s:                            # 1.500.000,50
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:                                       # 1500000,50
        s = s.replace(",", ".")
    elif s.count(".") > 1:                               # 1.500.000
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _add_meses(periodo: str, k: int) -> str:
    y, m = int(periodo[:4]), int(periodo[5:7])
    idx = y * 12 + (m - 1) + k
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


class FilaImport(BaseModel):
    cliente_id: int
    rol: str = ""
    forma: str = "Fa"
    periodo: str = ""
    meses: int = 1
    valor: str | None = None
    vendedor: str = ""
    pct_comision: str | None = None


class ConfirmarImportIn(BaseModel):
    filas: list[FilaImport]


def _leer_xlsx(raw: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    wb.close()
    return filas


@router.post("/servicios/importar")
async def importar_facturacion(file: UploadFile = File(...),
                               session: Session = Depends(get_session),
                               _u: Usuario = Depends(require_perm("clientes.importar"))):
    """Lee el Excel, cruza por NIT y devuelve una VISTA PREVIA (no guarda).
    Columnas (encabezado flexible): nit, nombre, rol, forma, periodo, valor,
    vendedor, comision, meses. Identifica clientes por NIT y roles por nombre."""
    raw = await file.read()
    try:
        filas_xl = _leer_xlsx(raw)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"No se pudo leer el Excel: {str(exc)[:120]}")
    if not filas_xl:
        raise HTTPException(400, "El Excel está vacío")

    # localizar la fila de encabezados (la primera que mapee al menos nit+periodo)
    hdr_idx, mapa = -1, {}
    for i, fila in enumerate(filas_xl[:15]):
        m = _detectar_columnas(list(fila))
        if "nit" in m and ("periodo" in m or "valor" in m):
            hdr_idx, mapa = i, m
            break
    if hdr_idx < 0:
        raise HTTPException(400, "No encontré la fila de encabezados (se requiere al "
                                 "menos 'nit' y 'periodo' o 'valor').")

    faltan = [c for c in ("nit", "periodo", "valor") if c not in mapa]
    if faltan:
        raise HTTPException(400, f"Faltan columnas obligatorias: {', '.join(faltan)}")

    # índices por NIT (solo dígitos) y roles existentes (por nombre normalizado)
    clientes = {_solo_dig(c.num_id): c for c in session.exec(select(Cliente)).all()}
    roles_norm = {_norm_hdr(r.descripcion): r for r in session.exec(select(RolServicio)).all()}

    def cel(fila, clave):
        idx = mapa.get(clave)
        return fila[idx] if (idx is not None and idx < len(fila)) else None

    ok, errores = [], []
    roles_nuevos: set[str] = set()
    n_cobros = 0
    for i in range(hdr_idx + 1, len(filas_xl)):
        fila = filas_xl[i]
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            continue
        nfila = i + 1  # 1-based para el usuario
        nit_raw = cel(fila, "nit")
        nit = _solo_dig(nit_raw)
        rol_txt = str(cel(fila, "rol") or "").strip()
        periodo = _norm_periodo(cel(fila, "periodo"))
        valor = _money(cel(fila, "valor"))
        forma = _norm_forma(cel(fila, "forma"))
        try:
            meses = int(float(cel(fila, "meses") or 1))
        except (ValueError, TypeError):
            meses = 1
        meses = max(1, min(meses, 60))
        vendedor = str(cel(fila, "vendedor") or "").strip()
        pct = _money(cel(fila, "pct"))

        cli = clientes.get(nit) if nit else None
        motivos = []
        if not nit:
            motivos.append("NIT vacío")
        elif not cli:
            motivos.append("NIT no encontrado")
        if not periodo:
            motivos.append("período inválido")
        if valor is None:
            motivos.append("valor inválido")

        if motivos:
            errores.append({"fila": nfila, "nit": str(nit_raw or ""),
                            "nombre": str(cel(fila, "nombre") or ""),
                            "rol": rol_txt, "periodo": periodo,
                            "motivo": ", ".join(motivos)})
            continue

        rol_existe = _norm_hdr(rol_txt) in roles_norm if rol_txt else True
        if rol_txt and not rol_existe:
            roles_nuevos.add(rol_txt)
        n_cobros += meses
        ok.append({
            "fila": nfila, "cliente_id": cli.id, "nit": cli.num_id,
            "nombre": cli.nombre, "rol": rol_txt, "rol_nuevo": bool(rol_txt and not rol_existe),
            "forma": forma, "periodo": periodo,
            "periodo_fin": _add_meses(periodo, meses - 1), "meses": meses,
            "valor": str(valor), "vendedor": vendedor,
            "pct_comision": (str(pct) if pct is not None else ""),
        })

    return {
        "resumen": {
            "filas_ok": len(ok), "filas_error": len(errores),
            "cobros_a_crear": n_cobros, "roles_nuevos": sorted(roles_nuevos),
        },
        "ok": ok, "errores": errores,
    }


@router.post("/servicios/importar/confirmar")
def confirmar_importacion(data: ConfirmarImportIn, session: Session = Depends(get_session),
                          _u: Usuario = Depends(require_perm("clientes.importar"))):
    """Crea los cobros programados (ClienteServicio) de las filas validadas.
    Expande ``meses`` en períodos consecutivos. Evita duplicar un cobro ya
    existente (mismo cliente + rol + período). Crea los roles que no existan."""
    roles_norm = {_norm_hdr(r.descripcion): r for r in session.exec(select(RolServicio)).all()}

    def resolver_rol(txt: str):
        txt = (txt or "").strip()
        if not txt:
            return None
        r = roles_norm.get(_norm_hdr(txt))
        if r:
            return r.id
        nuevo = RolServicio(descripcion=txt, activo=True)
        session.add(nuevo)
        session.flush()
        roles_norm[_norm_hdr(txt)] = nuevo
        return nuevo.id

    # cobros ya existentes por (cliente, rol, periodo) para no duplicar
    existentes = {(s.cliente_id, s.rol_servicio_id, s.periodo)
                  for s in session.exec(select(ClienteServicio)).all()}

    creados, omitidos = 0, 0
    for f in data.filas:
        rol_id = resolver_rol(f.rol)
        valor = _money(f.valor)
        pct = _money(f.pct_comision)
        meses = max(1, min(int(f.meses or 1), 60))
        for k in range(meses):
            per = _add_meses(f.periodo, k)
            clave = (f.cliente_id, rol_id, per)
            if clave in existentes:
                omitidos += 1
                continue
            session.add(ClienteServicio(
                cliente_id=f.cliente_id, rol_servicio_id=rol_id,
                forma_facturacion=f.forma or "Fa", periodo=per,
                valor_servicio=valor, vendedor=f.vendedor or "",
                pct_comision=pct))
            existentes.add(clave)
            creados += 1
    session.commit()
    return {"creados": creados, "omitidos": omitidos}
